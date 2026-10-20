"""
Karma (hibrit) Word çıktısı — taranmış/tasarımlı sayfa için.

  * Arka plan: sayfa görüntüsü, ÜZERİNDEKİ YAZILAR SİLİNMİŞ halde (fotoğraf, panel, şerit,
    çizgi, simge ve madde işaretleri kalır).
  * Yazılar: gerçek, düzenlenebilir Word metni; her satır orijinaldeki yerine sabitlenir.
    Punto ve harf aralığı, orijinal satır genişliğini tutturacak şekilde hesaplanır.
  * İsteğe bağlı: Claude görüntüye bakarak OCR yazım hatalarını düzeltir (ücretli; kullanım ölçülür).
"""
from __future__ import annotations

import base64
import io
import json
import urllib.request
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
from PIL import Image

from src import pdf_designed_scan as ds

_Y_OFFSET = 0.2  # satır kutusunun üstü ile harf üstü arasındaki fark (punto oranı)


def _claude_fix(img: Image.Image, lines: List[Dict[str, Any]], api_key: str, model: str) -> Dict[str, Any]:
    """OCR satırlarını görüntüye bakarak düzelt. {'usage':..., 'fixed': {id: text}}"""
    im = img.copy()
    im.thumbnail((1500, 2100))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=85)
    payload = [{"id": i, "text": l["text"]} for i, l in enumerate(lines)]
    prompt = (
        "Ekteki görüntü bir belge sayfası. Aşağıda OCR ile okunmuş satırlar var (JSON). "
        "Görüntüye bakarak SADECE OCR yazım/harf hatalarını düzelt (Türkçe karakterler, rakamlar, e-posta). "
        "Satır ekleme, silme, birleştirme, yeniden yazma YAPMA; görüntüde yazmayan hiçbir şeyi uydurma. "
        "Yalnızca değişen satırları şu biçimde döndür, başka hiçbir şey yazma: "
        '{"fixes": [{"id": 3, "text": "düzeltilmiş satır"}]}\n\n' + json.dumps(payload, ensure_ascii=False)
    )
    body = {
        "model": model, "max_tokens": 4000,
        "messages": [{"role": "user", "content": [
            {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                          "data": base64.b64encode(buf.getvalue()).decode()}},
            {"type": "text", "text": prompt}]}],
    }
    req = urllib.request.Request(
        "https://api.anthropic.com/v1/messages", data=json.dumps(body).encode(), method="POST",
        headers={"x-api-key": api_key, "anthropic-version": "2023-06-01", "content-type": "application/json"})
    with urllib.request.urlopen(req, timeout=120) as r:
        res = json.loads(r.read().decode())
    text = "".join(c.get("text", "") for c in res.get("content", []) if c.get("type") == "text")
    s, e = text.find("{"), text.rfind("}")
    fixed: Dict[int, str] = {}
    try:
        for f in json.loads(text[s:e + 1]).get("fixes", []):
            fixed[int(f["id"])] = str(f["text"]).strip()
    except Exception:
        pass
    return {"usage": res.get("usage", {}), "fixed": fixed}


def _erase_text(arr: np.ndarray, lines: List[Dict[str, Any]]) -> np.ndarray:
    """Yalnızca satır kutularındaki YAZI piksellerini sil; çevre zemine dokunma (gölge/leke kalmasın)."""
    out = arr.copy()
    mask = np.zeros(arr.shape[:2], dtype=np.uint8)
    for l in lines:
        y0, y1 = max(0, l["top"] - 4), min(arr.shape[0], l["bottom"] + 5)
        x0, x1 = max(0, l["left"] - 4), min(arr.shape[1], l["right"] + 5)
        if y1 <= y0 or x1 <= x0:
            continue
        reg = arr[y0:y1, x0:x1, :].astype(np.int32)
        bg = np.array(l.get("local_bg") or ds._dominant(reg.reshape(-1, 3))[0], dtype=np.int32)
        d = np.sqrt(((reg - bg) ** 2).sum(axis=2))
        mask[y0:y1, x0:x1] |= (d > 38).astype(np.uint8)
    try:
        import cv2
        mask = cv2.dilate(mask, np.ones((5, 5), np.uint8))
        return cv2.inpaint(out, mask * 255, 3, cv2.INPAINT_TELEA)
    except Exception:
        for l in lines:
            y0, y1 = max(0, l["top"] - 3), min(out.shape[0], l["bottom"] + 4)
            x0, x1 = max(0, l["left"] - 3), min(out.shape[1], l["right"] + 4)
            out[y0:y1, x0:x1, :] = np.array(l.get("zbg") or l.get("local_bg") or (255, 255, 255), dtype=np.uint8)
        return out


def _anchor_background(doc, png_bytes: bytes, w_pt: float, h_pt: float) -> None:
    """Görüntüyü sayfaya sabit, metnin ARKASINDA duran şekilde ekle."""
    from docx.shared import Pt
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    p = doc.add_paragraph()
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    shape = run.add_picture(io.BytesIO(png_bytes), width=Pt(w_pt), height=Pt(h_pt))
    inline = shape._inline
    graphic = inline.graphic
    cx, cy = inline.extent.cx, inline.extent.cy
    docpr_id = inline.docPr.id
    anchor = parse_xml(
        '<wp:anchor %s distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" '
        'behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="%d" cy="%d"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:docPr id="%d" name="Arka plan"/><wp:cNvGraphicFramePr/></wp:anchor>'
        % (nsdecls("wp"), cx, cy, docpr_id))
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


def _runs(l: Dict[str, Any], text: str) -> List[Tuple[str, bool]]:
    words = text.split(" ")
    flags = l.get("bold_flags") or []
    runs: List[Tuple[str, bool]] = []
    for i, w in enumerate(words):
        bold = flags[i] if len(words) == len(flags) and i < len(flags) else False
        if runs and runs[-1][1] == bold:
            runs[-1] = (runs[-1][0] + " " + w, bold)
        else:
            runs.append(((" " if runs else "") + w, bold))
    return runs


def _natural_width(runs: List[Tuple[str, bool]], size: float) -> float:
    import fitz
    # Yerleşik Helvetica yalnızca Latin-1 biliyor; Türkçe harfler 0 genişlik sayılıp punto şişmesin
    tr = str.maketrans("İıŞşĞğ", "IiSsGg")
    return sum(fitz.get_text_length(t.translate(tr), fontname="hebo" if b else "helv", fontsize=size)
               for t, b in runs)


def _fit_sizes(lines: List[Dict[str, Any]], texts: List[str], k: float) -> List[float]:
    """Her satırın punto değeri: orijinal satır genişliğini Arial-benzeri yazıyla tutturan değer.
    OCR gürültüsü punto oynatmasın diye benzer değerler tek seviyeye toplanır; seviye hiçbir zaman
    satırın kendi genişliğinden büyük olmaz (satır taşıp bölünmesin)."""
    raw: List[Optional[float]] = []
    for l, t in zip(lines, texts):
        w = (l["right"] - l["left"]) * k
        nat = _natural_width(_runs(l, t), 1.0)
        raw.append(w / nat if nat > 0 and len(t.replace(" ", "")) >= 6 else None)
    good = sorted(v for v in raw if v)
    levels: List[float] = []
    cur: List[float] = []
    for v in good:
        if cur and v > cur[0] * 1.08:
            levels.append(sorted(cur)[len(cur) // 2])
            cur = []
        cur.append(v)
    if cur:
        levels.append(sorted(cur)[len(cur) // 2])
    out: List[float] = []
    for l, r in zip(lines, raw):
        if r is None:
            r = ds._font_pt(l["h"], k)
        cand = [v for v in levels if r / 1.08 <= v <= r * 1.02]
        if cand:
            r = max(cand)
        out.append(max(6.0, min(40.0, round(r * 2) / 2.0)))
    return out


def _frame_line(doc, l: Dict[str, Any], k: float, size: float, text: str, page_w: float = 0.0) -> None:
    from docx.shared import Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    target = (l["right"] - l["left"]) * k
    fr = OxmlElement("w:framePr")
    wid = min(target * 1.35 + 84, page_w - l["left"] * k - 2) if page_w else target * 1.35 + 84
    fr.set(qn("w:w"), str(int(wid * 20)))
    fr.set(qn("w:wrap"), "around")
    fr.set(qn("w:vAnchor"), "page")
    fr.set(qn("w:hAnchor"), "page")
    fr.set(qn("w:x"), str(int(max(0, l["left"] * k) * 20)))
    fr.set(qn("w:y"), str(int(max(0, l["top"] * k - _Y_OFFSET * size) * 20)))
    p._p.get_or_add_pPr().insert(0, fr)

    runs = _runs(l, text)
    nchar = max(1, sum(len(t) for t, _ in runs))
    spacing = (target - _natural_width(runs, size)) / nchar
    spacing = max(-0.5, min(0.8, spacing))
    for t, b in runs:
        r = p.add_run(t)
        r.font.size = Pt(size)
        r.font.name = "Arial"
        r.bold = b
        r.font.color.rgb = RGBColor(*l["color"])
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(round(spacing * 20))))
        r._element.get_or_add_rPr().append(sp)


def hybrid_pdf_to_word(pdf_path: str, docx_path: str, *, use_claude: bool = False,
                       api_key: str = "", model: str = "claude-haiku-5-5") -> Dict[str, Any]:
    import fitz
    from docx import Document
    from docx.shared import Pt

    info: Dict[str, Any] = {"usage": {}, "fixes": 0}
    doc = fitz.open(pdf_path)
    dw = Document()
    sec = dw.sections[0]
    wpt, hpt = doc.load_page(0).rect.width, doc.load_page(0).rect.height
    sec.page_width, sec.page_height = Pt(wpt), Pt(hpt)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(0)

    for i in range(len(doc)):
        pg = doc.load_page(i)
        scale = max(1.5, min(3.0, 1700.0 / max(1.0, pg.rect.width)))
        pix = pg.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False, colorspace=fitz.csRGB)
        img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
        arr = np.asarray(img).copy()
        lay = ds.analyze_page(img) or {"zones": [{"x0": 0, "x1": img.width, "role": "main", "bands": [],
                                                  "bg": ds._dominant(arr.reshape(-1, 3))[0]}]}
        lines: List[Dict[str, Any]] = []
        for z in lay["zones"]:
            photos = ds._find_photos(arr, z["bg"], z["x0"], z["x1"])
            for it in ds._build_items(arr, z, photos):
                if it["kind"] == "text":
                    for l in it["lines"]:
                        mid = (l["top"] + l["bottom"]) / 2
                        bd = next((b for b in z["bands"] if b["y0"] - 4 <= mid <= b["y1"] + 4), None)
                        l["zbg"] = bd["color"] if bd else z["bg"]
                    lines.extend(it["lines"])
        k = wpt / float(img.width)

        texts = [l["text"] for l in lines]
        if use_claude and lines:
            res = _claude_fix(img, lines, api_key, model)
            for kk, v in res["usage"].items():
                if isinstance(v, int):
                    info["usage"][kk] = info["usage"].get(kk, 0) + v
            for idx, t in res["fixed"].items():
                if 0 <= idx < len(texts) and t:
                    if len(t.split(" ")) != len(texts[idx].split(" ")):
                        lines[idx]["bold_flags"] = []
                    texts[idx] = t
                    info["fixes"] += 1

        clean = Image.fromarray(_erase_text(arr, lines))
        buf = io.BytesIO()
        clean.save(buf, "PNG")
        if i > 0:
            dw.add_page_break()
        _anchor_background(dw, buf.getvalue(), wpt, hpt)
        sizes = _fit_sizes(lines, texts, k)
        for l, t, sz in zip(lines, texts, sizes):
            _frame_line(dw, l, k, sz, t, wpt)
    dw.save(docx_path)
    doc.close()
    return info
