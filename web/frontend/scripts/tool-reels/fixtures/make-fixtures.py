"""
Araç videoları için örnek belgeler (hepsi uydurma, gerçek veri yok).

  Faaliyet-Raporu-2025.pdf  → PDF Sıkıştır (görselli, ~8 MB; "beklenen küçülme" çıksın diye)
  Sozlesme/Fatura/Ekler.pdf → PDF Birleştir

Kullanım: python make-fixtures.py <çıktı klasörü> <yazı tipi klasörü>
"""
import fitz, io, math, random, sys
from PIL import Image, ImageDraw, ImageFilter, ImageFont
random.seed(11)
S = sys.argv[1]  # çıktı klasörü
FONTS = sys.argv[2]  # Roboto-Bold.ttf / Roboto-Regular.ttf klasörü (public/fonts)
RB, RR = FONTS + "/Roboto-Bold.ttf", FONTS + "/Roboto-Regular.ttf"
W, H = 2000, 1250


def noise(img, amt=0.10):
    n = Image.effect_noise(img.size, 70).convert("RGB")
    return Image.blend(img, n, amt)


def grad(w, h, top, bot):
    g = Image.linear_gradient("L").resize((w, h))
    a, b = Image.new("RGB", (w, h), top), Image.new("RGB", (w, h), bot)
    return Image.composite(b, a, g)


def skyline():
    im = grad(W, H, (250, 140, 80), (30, 40, 90))
    d = ImageDraw.Draw(im)
    d.ellipse((1350, 420, 1650, 720), fill=(255, 220, 150))
    x = 0
    while x < W:
        bw, bh = random.randint(110, 240), random.randint(300, 900)
        col = random.choice([(22, 28, 52), (28, 34, 62), (18, 22, 44)])
        d.rectangle((x, H - bh, x + bw, H), fill=col)
        for wy in range(H - bh + 30, H - 30, 46):
            for wx in range(x + 18, x + bw - 18, 36):
                if random.random() < 0.42:
                    d.rectangle((wx, wy, wx + 14, wy + 22), fill=(255, 214, 120))
        x += bw + random.randint(0, 14)
    return noise(im.filter(ImageFilter.GaussianBlur(1.1)))


def mountains(top=(120, 170, 235), bot=(235, 225, 205)):
    im = grad(W, H, top, bot)
    d = ImageDraw.Draw(im)
    for layer, col in enumerate([(98, 120, 150), (70, 98, 120), (44, 70, 88), (24, 44, 56)]):
        base = 520 + layer * 170
        ph = random.uniform(0, 6)
        pts = [(0, H)] + [
            (x, base - 150 * math.sin(x / 260 + ph + layer) - 60 * math.sin(x / 90 + layer * 2))
            for x in range(0, W + 20, 20)
        ] + [(W, H)]
        d.polygon(pts, fill=col)
    return noise(im.filter(ImageFilter.GaussianBlur(1.0)))


def abstract(c1, c2, c3):
    im = grad(W, H, c1, c2)
    d = ImageDraw.Draw(im, "RGBA")
    for _ in range(14):
        r = random.randint(120, 420)
        x, y = random.randint(0, W), random.randint(0, H)
        d.ellipse((x - r, y - r, x + r, y + r), fill=c3 + (random.randint(40, 110),))
    return noise(im.filter(ImageFilter.GaussianBlur(6)), 0.12)


def font(p, s):
    return ImageFont.truetype(p, s)


def bar_chart():
    im = Image.new("RGB", (W, 1100), (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.text((60, 30), "Çeyreklik ciro (milyon ₺)", font=font(RB, 56), fill=(20, 30, 60))
    for g in range(0, 5):
        y = 960 - g * 200
        d.line((160, y, W - 100, y), fill=(225, 230, 240), width=3)
    for i, (lab, v) in enumerate([("Q1", 42), ("Q2", 58), ("Q3", 71), ("Q4", 96)]):
        x = 220 + i * 420
        hgt = v * 8.4
        d.rectangle((x, 960 - hgt, x + 300, 960), fill=(37, 99, 235) if i < 3 else (16, 185, 129))
        d.text((x + 150, 960 - hgt - 64), str(v), font=font(RB, 50), fill=(20, 30, 60), anchor="ma")
        d.text((x + 150, 985), lab, font=font(RR, 46), fill=(70, 80, 110), anchor="ma")
    return noise(im, 0.035)


def pie_chart():
    im = Image.new("RGB", (W, 1100), (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.text((60, 30), "Gelirin dağılımı", font=font(RB, 56), fill=(20, 30, 60))
    parts = [("Kara taşımacılığı", 46, (37, 99, 235)), ("Depolama", 24, (16, 185, 129)),
             ("Deniz yolu", 18, (245, 158, 11)), ("Diğer", 12, (148, 163, 184))]
    start = -90
    for name, p, col in parts:
        d.pieslice((120, 170, 900, 950), start, start + p * 3.6, fill=col)
        start += p * 3.6
    for i, (name, p, col) in enumerate(parts):
        y = 330 + i * 150
        d.rectangle((1100, y, 1160, y + 60), fill=col)
        d.text((1190, y - 6), f"{name}  %{p}", font=font(RR, 54), fill=(30, 40, 70))
    return noise(im, 0.035)


def jpg(im, q=95):
    b = io.BytesIO()
    im.save(b, "JPEG", quality=q)
    return b.getvalue()


doc = fitz.open()


def page():
    p = doc.new_page(width=595, height=842)
    p.insert_font("rb", RB)
    p.insert_font("rr", RR)
    return p


def T(p, xy, s, size, bold=False, color=(0.08, 0.12, 0.25)):
    p.insert_text(xy, s, fontname="rb" if bold else "rr", fontsize=size, color=color)


def para(p, rect, s, size=11.5):
    p.insert_textbox(fitz.Rect(*rect), s, fontname="rr", fontsize=size, color=(0.18, 0.2, 0.3), lineheight=1.45)


def heading(p, title, underline_to):
    T(p, (46, 70), title, 26, True)
    p.draw_line((46, 84), (underline_to, 84), color=(0.15, 0.4, 0.9), width=3)


lorem = [
    "2025 yılı, Atlas Lojistik için büyümenin ve dijital dönüşümün hızlandığı bir yıl oldu. Yeni depo yatırımlarımız ve rota optimizasyonu çalışmalarımız sayesinde teslimat sürelerini ortalama yüzde on dört kısalttık.",
    "Müşteri memnuniyeti puanımız beş üzerinden 4,7 seviyesine ulaştı. Anlık takip sistemimiz artık tüm sevkiyatların yüzde doksan sekizini kapsıyor ve müşterilerimiz yüklerini canlı olarak izleyebiliyor.",
    "Önümüzdeki dönemde Anadolu'daki bölgesel dağıtım ağımızı genişletmeyi, filomuzun bir bölümünü elektrikli araçlara dönüştürmeyi ve karbon salımımızı yüzde yirmi azaltmayı hedefliyoruz.",
]
IH = 503 * H / W  # tam genişlik görsel yüksekliği

# 1 Kapak
p = page()
p.draw_rect(fitz.Rect(0, 0, 595, 842), color=None, fill=(0.04, 0.07, 0.16))
p.insert_image(fitz.Rect(0, 300, 595, 300 + 595 * H / W), stream=jpg(skyline()))
T(p, (46, 120), "ATLAS LOJİSTİK", 14, True, (0.45, 0.65, 1.0))
T(p, (46, 175), "Yıllık Faaliyet", 40, True, (1, 1, 1))
T(p, (46, 222), "Raporu 2025", 40, True, (1, 1, 1))
T(p, (46, 262), "Büyüme, verimlilik ve sürdürülebilirlik", 14, False, (0.75, 0.82, 0.95))
# 2 Yönetim özeti
p = page()
heading(p, "Yönetim Özeti", 160)
para(p, (46, 105, 549, 330), "\n\n".join(lorem[:2]))
p.insert_image(fitz.Rect(46, 330, 549, 330 + IH), stream=jpg(mountains()))
para(p, (46, 330 + IH + 18, 549, 820), lorem[2])
# 3 Ciro
p = page()
heading(p, "Finansal Göstergeler", 230)
p.insert_image(fitz.Rect(46, 110, 549, 110 + 503 * 1100 / W), stream=jpg(bar_chart(), 93))
para(p, (46, 110 + 503 * 1100 / W + 16, 549, 820),
     "Dördüncü çeyrekte sezonluk talebin ve yeni müşteri sözleşmelerinin etkisiyle ciro bir önceki yıla göre yüzde otuz iki arttı. Yıl genelinde toplam ciro 267 milyon ₺ oldu.")
# 4 Tablo
p = page()
heading(p, "Operasyonel Veriler", 225)
rows = [("Gösterge", "2024", "2025", "Değişim"), ("Teslimat sayısı", "1.248.000", "1.612.000", "+%29"),
        ("Ort. teslimat süresi", "41 sa", "35 sa", "-%14"), ("Depo alanı (m²)", "86.000", "118.000", "+%37"),
        ("Filo (araç)", "412", "506", "+%23"), ("Müşteri puanı", "4,4", "4,7", "+0,3")]
y = 112
for i, r in enumerate(rows):
    if i == 0:
        p.draw_rect(fitz.Rect(46, y - 18, 549, y + 10), fill=(0.08, 0.15, 0.4), color=None)
    elif i % 2 == 0:
        p.draw_rect(fitz.Rect(46, y - 18, 549, y + 10), fill=(0.94, 0.96, 1), color=None)
    for j, c in enumerate(r):
        T(p, (56 + [0, 230, 330, 440][j], y), c, 11.5, i == 0, (1, 1, 1) if i == 0 else (0.1, 0.13, 0.25))
    y += 34
p.insert_image(fitz.Rect(46, 360, 549, 360 + IH), stream=jpg(abstract((20, 60, 140), (10, 20, 60), (80, 160, 255))))
para(p, (46, 360 + IH + 18, 549, 820),
     "Tüm göstergelerde iyileşme sağlandı. Depo alanındaki artış, yeni Gebze ve İzmir tesislerinin devreye alınmasından kaynaklanıyor.")
# 5 Pasta
p = page()
heading(p, "Gelir Kaynakları", 200)
p.insert_image(fitz.Rect(46, 110, 549, 110 + 503 * 1100 / W), stream=jpg(pie_chart(), 93))
sun = mountains((255, 190, 130), (255, 235, 200)).crop((0, 0, W, int(H * 0.62)))
p.insert_image(fitz.Rect(46, 110 + 503 * 1100 / W + 14, 549, 110 + 503 * 1100 / W + 14 + 503 * H / W * 0.62), stream=jpg(sun))
# 6 Galeri
p = page()
heading(p, "Tesislerimizden Kareler", 270)
cols = [((30, 90, 60), (10, 40, 30), (120, 220, 150)), ((120, 40, 40), (50, 15, 30), (255, 140, 120)),
        ((60, 60, 120), (20, 20, 60), (170, 150, 255)), ((20, 110, 120), (8, 45, 60), (110, 230, 230))]
names = ["Gebze Merkez Depo", "İzmir Aktarma Hattı", "Ankara Dağıtım Noktası", "Mersin Liman Terminali"]
for i, c in enumerate(cols):
    x, y = 46 + (i % 2) * 256, 112 + (i // 2) * 330
    h = 247 * H / W * 1.25
    p.insert_image(fitz.Rect(x, y, x + 247, y + h), stream=jpg(abstract(*c).resize((1500, 940))))
    para(p, (x, y + h + 4, x + 247, y + 320), names[i], 10.5)
# 7 Kapanış
p = page()
p.draw_rect(fitz.Rect(0, 0, 595, 842), color=None, fill=(0.04, 0.07, 0.16))
p.insert_image(fitz.Rect(0, 0, 595, 595 * H / W), stream=jpg(skyline()))
T(p, (46, 450), "Teşekkürler", 36, True, (1, 1, 1))
T(p, (46, 490), "İletişim: info@atlaslojistik.example · 0850 000 00 00", 12, False, (0.75, 0.82, 0.95))
doc.save(S + "/Faaliyet-Raporu-2025.pdf", deflate=True)
print(len(doc), "sayfa")


# ── PDF Birleştir için küçük belgeler ─────────────────────────────────────────
for name, title, pages in [("Sozlesme", "Hizmet Sözleşmesi", 3), ("Fatura", "Fatura No: 2026-0412", 2), ("Ekler", "Ekler ve Belgeler", 2)]:
    d = fitz.open()
    for i in range(pages):
        pg = d.new_page(width=595, height=842)
        pg.insert_text((60, 100), title, fontsize=26, fontname="helv")
        pg.insert_text((60, 150), "Sayfa %d" % (i + 1), fontsize=14, fontname="helv")
        for k in range(12):
            pg.insert_text((60, 210 + k * 24), "Örnek satır metni %d - demo belge" % (k + 1), fontsize=11, fontname="helv")
    d.save("%s/%s.pdf" % (S, name))
print("birleştirme belgeleri hazır")


# ══ Diğer araçlar için örnek belgeler (hepsi uydurma) ═════════════════════════
import zipfile, os

# 1) Hafif rapor: Faaliyet raporunun küçültülmüş hâli (sayfa/dönüştürme araçları hızlı çalışsın).
light = fitz.open(S + "/Faaliyet-Raporu-2025.pdf")
try:
    light.rewrite_images(dpi_threshold=110, dpi_target=80, quality=55)
except Exception as e:  # eski PyMuPDF: görselleri küçültemeden kaydet
    print("rewrite_images yok:", e)
light.save(S + "/Rapor-Hafif.pdf", garbage=4, deflate=True)
light.close()

# 2) Görseller (JPG): yüksek çözünürlüklü, sıkıştır/boyutlandır için.
jpg_names = [("Manzara.jpg", mountains), ("Sehir.jpg", skyline), ("Soyut.jpg", lambda: abstract((20, 80, 200), (140, 60, 220), (250, 170, 60)))]
for fname, fn in jpg_names:
    im = fn().resize((3200, 2000), Image.LANCZOS)
    with open(S + "/" + fname, "wb") as f:
        f.write(jpg(noise(im, 0.03), 96))

# 3) Taranmış belge: yalnız görselden oluşan, hafif eğik 2 sayfa (OCR / aranabilir PDF / eğri tarama).
def scan_page(idx, angle):
    im = Image.new("RGB", (1240, 1754), (246, 244, 238))
    d = ImageDraw.Draw(im)
    d.text((110, 140), "KİRA SÖZLEŞMESİ" if idx == 0 else "SÖZLEŞME EKLERİ", font=font(RB, 64), fill=(25, 25, 30))
    lines = [
        "Madde 1 - Taraflar: Kiraya veren ile kiracı aşağıdaki",
        "şartlarda anlaşmışlardır. Bu belge örnek amaçlıdır.",
        "Madde 2 - Kira bedeli aylık 18.500 TL olup her ayın",
        "beşinci gününe kadar ödenir.",
        "Madde 3 - Kira süresi bir yıldır; süre sonunda taraflar",
        "yazılı bildirimle sözleşmeyi yenileyebilir.",
        "Madde 4 - Depozito bedeli iki aylık kira tutarındadır.",
        "Madde 5 - Taşınmaz yalnızca konut olarak kullanılır.",
    ]
    for i, s in enumerate(lines):
        d.text((110, 300 + i * 78), s, font=font(RR, 46), fill=(30, 30, 40))
    im = noise(im, 0.07).filter(ImageFilter.GaussianBlur(0.8))
    return im.rotate(angle, resample=Image.BICUBIC, expand=False, fillcolor=(246, 244, 238))

sc = fitz.open()
for idx, ang in enumerate((1.8, -1.4)):
    pg = sc.new_page(width=595, height=842)
    pg.insert_image(pg.rect, stream=jpg(scan_page(idx, ang), 88))
sc.save(S + "/Taranmis-Belge.pdf", deflate=True)
sc.close()

# 4) Doldurulabilir form (AcroForm).
fm = fitz.open()
fp = fm.new_page(width=595, height=842)
fp.insert_text((60, 90), "Başvuru Formu", fontsize=26, fontname="helv")
fields = [("Ad Soyad", "ad"), ("E-posta", "eposta"), ("Telefon", "telefon"), ("Şehir", "sehir")]
for i, (lab, nm) in enumerate(fields):
    y = 150 + i * 70
    fp.insert_text((60, y), lab, fontsize=12, fontname="helv")
    w = fitz.Widget()
    w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    w.field_name = nm
    w.rect = fitz.Rect(60, y + 8, 400, y + 34)
    w.text_fontsize = 12
    fp.add_widget(w)
fp.insert_text((60, 450), "Bilgilendirme metnini okudum", fontsize=12, fontname="helv")
cb = fitz.Widget()
cb.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX
cb.field_name = "onay"
cb.rect = fitz.Rect(300, 438, 316, 454)
fp.add_widget(cb)
fm.save(S + "/Basvuru-Formu.pdf")
fm.close()

# 5) Şifreli PDF (parola: demo123) ve 6) bozuk PDF.
base = fitz.open(S + "/Sozlesme.pdf")
base.save(S + "/Sifreli-Sozlesme.pdf", encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="demo123", owner_pw="demo123-owner")
base.close()
raw = open(S + "/Rapor-Hafif.pdf", "rb").read()
open(S + "/Bozuk-Rapor.pdf", "wb").write(raw[: int(len(raw) * 0.82)])  # sonu kesik: açılmaz

# 7) Ön yüz / arka yüz (dönüşümlü birleştirme): tek sayfalar.
for name, nums in (("Tarama-Ön-Yuzler", (1, 3, 5)), ("Tarama-Arka-Yuzler", (6, 4, 2))):
    d = fitz.open()
    for n_ in nums:
        pg = d.new_page(width=595, height=842)
        pg.insert_text((60, 100), "Sayfa %d" % n_, fontsize=48, fontname="helv")
        pg.insert_text((60, 160), "Çift taraflı tarama örneği", fontsize=16, fontname="helv")
    d.save("%s/%s.pdf" % (S, name.replace("ö", "o").replace("Ö", "O")))

# 8) Tablolu PDF (Excel'e aktar).
tb = fitz.open()
tp = tb.new_page(width=595, height=842)
tp.insert_text((60, 90), "Satış Özeti - 2026", fontsize=22, fontname="helv")
rows = [("Ürün", "Adet", "Birim fiyat", "Tutar"), ("Klavye", "40", "750", "30.000"), ("Fare", "85", "320", "27.200"), ("Monitör", "22", "6.400", "140.800"), ("Kulaklık", "60", "980", "58.800"), ("Dock", "18", "2.150", "38.700")]
for r, row in enumerate(rows):
    for c, cell in enumerate(row):
        x = 60 + c * 125
        y = 140 + r * 34
        tp.draw_rect(fitz.Rect(x, y, x + 125, y + 34), color=(0.55, 0.6, 0.7), fill=(0.9, 0.94, 1) if r == 0 else None, width=0.8)
        tp.insert_text((x + 8, y + 22), cell, fontsize=12, fontname="helv")
tb.save(S + "/Satis-Tablosu.pdf")
tb.close()

# 9) Office belgeleri.
from docx import Document
dx = Document()
dx.add_heading("Hizmet Teklifi", 0)
dx.add_paragraph("Sayın yetkili, talep ettiğiniz web sitesi yenileme çalışması için teklifimiz aşağıdadır.")
t_ = dx.add_table(rows=3, cols=3)
for r, row in enumerate((("Kalem", "Süre", "Tutar"), ("Tasarım", "2 hafta", "45.000 TL"), ("Geliştirme", "4 hafta", "90.000 TL"))):
    for c, v in enumerate(row):
        t_.cell(r, c).text = v
dx.add_paragraph("Teklif 30 gün geçerlidir.")
dx.save(S + "/Teklif.docx")

from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Satışlar"
for row in (("Ay", "Gelir", "Gider"), ("Ocak", 120000, 80000), ("Şubat", 135000, 82000), ("Mart", 150000, 91000), ("Nisan", 142000, 88000)):
    ws.append(row)
wb.save(S + "/Satis-Tablosu.xlsx")

from pptx import Presentation
pr = Presentation()
s1 = pr.slides.add_slide(pr.slide_layouts[0])
s1.shapes.title.text = "2026 Planı"
s1.placeholders[1].text = "Yıllık hedefler ve takvim"
s2 = pr.slides.add_slide(pr.slide_layouts[1])
s2.shapes.title.text = "Öncelikler"
s2.placeholders[1].text = "Müşteri memnuniyeti\nYeni pazarlar\nMaliyet kontrolü"
pr.save(S + "/Sunum.pptx")

# 10) UYAP UDF (örnek dilekçe).
text = "Sayın Hâkimliğe,Müvekkilim adına iş bu dilekçeyi sunarım."
xml = ('<?xml version="1.0" encoding="UTF-8" ?>\n<template format_id="1.8">\n<content><![CDATA[%s]]></content>\n'
       '<properties><pageFormat leftMargin="42.5" rightMargin="28.3" topMargin="14.1" bottomMargin="14.1" paperOrientation="1" /></properties>\n'
       '<elements resolver="hvl-default">\n<paragraph Alignment="1" LeftIndent="0.0" RightIndent="0.0"><content startOffset="0" length="16" family="Times New Roman" size="12" bold="true" /></paragraph>\n'
       '<paragraph Alignment="3" LeftIndent="0.0" RightIndent="0.0"><content startOffset="16" length="41" family="Times New Roman" size="12" /></paragraph>\n</elements>\n'
       '<styles><style name="hvl-default" family="Times New Roman" size="12" description="Gövde" /></styles>\n</template>') % text
with zipfile.ZipFile(S + "/Dilekce.udf", "w") as z:
    z.writestr("content.xml", xml)
print("diğer araç belgeleri hazır")
