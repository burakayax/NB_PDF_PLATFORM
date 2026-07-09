"""Araç "smoke" testi — her non-AI aracın çekirdek fonksiyonunu GERÇEK bir dosyayla
çalıştırır ve geçerli çıktı üretip üretmediğini doğrular.

Amaç: kod regresyonlarını (ör. pypdf tembel-okuma "seek of closed file" hatası)
deploy'dan ÖNCE yakalamak. Unit testler mock kullandığı için bunları göremiyordu.

Kural:
  * Saf-Python araçlar (fitz/pikepdf/pypdf/reportlab) → her ortamda çalışmalı; hata = FAIL.
  * Harici araç gerektirenler (LibreOffice / poppler / wkhtmltopdf) → araç yoksa SKIP,
    varsa çalışmalı. Böylece CI (ve prod imajı) neyin gerçekten kurulu olduğunu yansıtır.
"""

from __future__ import annotations

import os
import sys
import shutil
from pathlib import Path

import pytest

# Proje kökünü sys.path'e ekle (web/backend/tests/pdf/ → proje kökü 4 seviye üst).
_ROOT = Path(__file__).resolve().parents[4]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

import fitz  # PyMuPDF
from PIL import Image

from src import pdf_toolkit_extra as ptx  # noqa: E402

try:
    import src.pdf_engine as engine  # noqa: E402
    ENGINE_OK = True
    ENGINE_ERR = ""
except Exception as exc:  # pragma: no cover - CI'da tesseract kurulu
    engine = None  # type: ignore[assignment]
    ENGINE_OK = False
    ENGINE_ERR = str(exc)

needs_engine = pytest.mark.skipif(not ENGINE_OK, reason=f"pdf_engine import edilemedi: {ENGINE_ERR}")

# Hata mesajında bunlardan biri geçerse → harici araç eksikliği → SKIP (kod hatası değil).
_BINARY_HINTS = (
    "libreoffice", "soffice", "powerpoint", "microsoft word",
    "yalnızca windows", "only windows", "wkhtmltopdf", "poppler", "pdftoppm",
)


def _run(fn):
    """Aracı çalıştırır; harici-araç hatasında testi SKIP eder, gerçek hatada yükseltir."""
    try:
        return fn()
    except Exception as e:  # noqa: BLE001
        if any(h in str(e).lower() for h in _BINARY_HINTS):
            pytest.skip(f"Harici araç yok (prod imajında olmalı): {str(e)[:120]}")
        raise


def _valid_pdf(path: str, password: str | None = None) -> bool:
    if not os.path.exists(path) or os.path.getsize(path) < 100:
        return False
    try:
        d = fitz.open(path)
        if password is not None:
            d.authenticate(password)
        n = d.page_count
        d.close()
        return n > 0
    except Exception:
        return False


def _nonempty(path: str, minsize: int = 50) -> bool:
    return os.path.exists(path) and os.path.getsize(path) > minsize


@pytest.fixture(scope="module")
def assets(tmp_path_factory):
    d = tmp_path_factory.mktemp("tools_smoke")
    pdf = d / "in.pdf"
    doc = fitz.open()
    for i in range(3):
        pg = doc.new_page()
        pg.insert_text((72, 120), f"Smoke test sayfa {i + 1} - ornek metin ABC 123", fontsize=16)
    doc.save(str(pdf))
    doc.close()

    enc = d / "enc.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 120), "gizli icerik", fontsize=16)
    doc.save(str(enc), encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="1234", owner_pw="1234")
    doc.close()

    png = d / "img.png"
    Image.new("RGB", (400, 300), (200, 180, 140)).save(str(png))

    from docx import Document
    docx = d / "in.docx"
    doc2 = Document()
    doc2.add_paragraph("Smoke test Word belgesi - ABC 123")
    doc2.save(str(docx))

    from openpyxl import Workbook
    xlsx = d / "in.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Ad", "Tutar"])
    ws.append(["Kalem", 250])
    wb.save(str(xlsx))

    return {
        "dir": d, "pdf": str(pdf), "enc": str(enc),
        "png": str(png), "docx": str(docx), "xlsx": str(xlsx),
    }


def _out(assets, name: str) -> str:
    return str(assets["dir"] / name)


# ── Saf-Python (her ortamda çalışmalı) ──────────────────────────────────────

def test_delete_pages(assets):
    o = _out(assets, "del.pdf")
    _run(lambda: ptx.delete_pages_pdf(assets["pdf"], o, [2]))
    assert _valid_pdf(o)


def test_rotate(assets):
    o = _out(assets, "rot.pdf")
    _run(lambda: ptx.rotate_pdf(assets["pdf"], o, 0, None, per_page_degrees={1: 90, 2: 180}))
    assert _valid_pdf(o)


def test_organize(assets):
    o = _out(assets, "org.pdf")
    _run(lambda: ptx.organize_pdf(assets["pdf"], o, [3, 1, 2]))
    assert _valid_pdf(o)


def test_watermark(assets):
    o = _out(assets, "wm.pdf")
    _run(lambda: ptx.add_watermark_text(assets["pdf"], o, "GIZLI", font_color="#FF0000"))
    assert _valid_pdf(o)


def test_page_numbers(assets):
    o = _out(assets, "pn.pdf")
    _run(lambda: ptx.add_page_numbers(assets["pdf"], o, start_at=1, position="bottom-center"))
    assert _valid_pdf(o)


def test_repair(assets):
    o = _out(assets, "rep.pdf")
    _run(lambda: ptx.repair_pdf(assets["pdf"], o))
    assert _valid_pdf(o)


def test_pdf_to_text(assets):
    o = _out(assets, "t.txt")
    _run(lambda: ptx.pdf_to_text(assets["pdf"], o))
    assert _nonempty(o)


def test_flatten(assets):
    o = _out(assets, "flat.pdf")
    _run(lambda: ptx.flatten_pdf(assets["pdf"], o))
    assert _valid_pdf(o)


def test_image_to_pdf(assets):
    o = _out(assets, "i2p.pdf")
    _run(lambda: ptx.images_to_pdf([assets["png"]], o))
    assert _valid_pdf(o)


def test_unlock(assets):
    o = _out(assets, "unl.pdf")
    _run(lambda: ptx.unlock_pdf_pikepdf(assets["enc"], o, "1234"))
    assert _valid_pdf(o)


def test_pdf_to_ppt(assets):
    o = _out(assets, "p2p.pptx")
    _run(lambda: ptx.pdf_to_pptx(assets["pdf"], o))
    assert _nonempty(o)


def test_html_to_pdf(assets):
    o = _out(assets, "h.pdf")
    _run(lambda: ptx.html_to_pdf_file("<html><body><h1>QA</h1><p>test 123</p></body></html>", o))
    assert _valid_pdf(o)


# ── pdf_engine tabanlı (tesseract gerekir → CI'da kurulu) ───────────────────

@needs_engine
def test_split(assets):
    o = _out(assets, "s.pdf")
    _run(lambda: engine.extract_pages(assets["pdf"], [1, 3], o))
    assert _valid_pdf(o)


@needs_engine
def test_merge(assets):
    o = _out(assets, "m.pdf")
    _run(lambda: engine.merge_pdfs([assets["pdf"], assets["pdf"]], o))
    assert _valid_pdf(o)


@needs_engine
def test_compress(assets):
    o = _out(assets, "c.pdf")
    _run(lambda: engine.compress_pdf(assets["pdf"], o, quality="auto"))
    assert _valid_pdf(o)


@needs_engine
def test_encrypt(assets):
    o = _out(assets, "e.pdf")
    _run(lambda: engine.encrypt_pdf(assets["pdf"], o, user_password="1234"))
    assert _valid_pdf(o, password="1234")


@needs_engine
def test_get_num_pages(assets):
    # REGRESYON KORUMASI: pypdf kapalı-dosya hatası ("seek of closed file").
    assert _run(lambda: engine.get_num_pages(assets["pdf"])) == 3


@needs_engine
def test_pdf_to_excel(assets):
    # REGRESYON KORUMASI: pypdf kapalı-dosya hatası bu aracı tamamen kırmıştı.
    o = _out(assets, "p2x.xlsx")
    _run(lambda: engine.pdf_text_to_excel(assets["pdf"], o))
    assert _nonempty(o)


@needs_engine
def test_classify_password_requirement(assets):
    # REGRESYON KORUMASI: şifre-gerekli tespiti de aynı hatadan etkileniyordu.
    req_enc, _ = _run(lambda: engine.classify_pdf_password_requirement(assets["enc"]))
    req_plain, _ = _run(lambda: engine.classify_pdf_password_requirement(assets["pdf"]))
    assert req_enc is True and req_plain is False


@needs_engine
def test_pdf_to_word(assets):
    o = _out(assets, "p2w.docx")
    _run(lambda: engine.pdf_to_word(assets["pdf"], o))
    assert _nonempty(o)


@needs_engine
def test_excel_to_pdf(assets):
    # reportlab yedeği ile LibreOffice olmadan da çalışmalı.
    o = _out(assets, "x2p.pdf")
    _run(lambda: engine.excel_to_pdf(assets["xlsx"], o))
    assert _valid_pdf(o)


@needs_engine
def test_word_to_pdf(assets):
    # LibreOffice/Word gerekir → yoksa _run SKIP eder.
    o = _out(assets, "w2p.pdf")
    _run(lambda: engine.word_to_pdf(assets["docx"], o))
    assert _valid_pdf(o)


# ── Harici ikili gerektiren rasterizasyon/ofis ──────────────────────────────

def test_pdf_to_image(assets):
    # poppler (pdftoppm) gerekir → yoksa SKIP.
    if not (shutil.which("pdftoppm") or shutil.which("pdftocairo")):
        pytest.skip("poppler (pdftoppm) yok")
    zip_path = _run(lambda: ptx.pdf_to_images_zip(assets["pdf"], str(assets["dir"])))
    assert _nonempty(zip_path)


def test_ppt_to_pdf(assets):
    # LibreOffice gerekir → yoksa SKIP.
    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        pytest.skip("LibreOffice (soffice) yok")
    src_pptx = _out(assets, "src.pptx")
    ptx.pdf_to_pptx(assets["pdf"], src_pptx)
    o = _out(assets, "ptp.pdf")
    _run(lambda: ptx.pptx_to_pdf(src_pptx, o))
    assert _valid_pdf(o)
